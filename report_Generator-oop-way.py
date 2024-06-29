from docx import Document
from openpyxl import Workbook
import os

# massage the list data and store them in a dictionary for easier retrieval
def massageListToDict(inputList):
    #remove empty record from the list
    inputList = list(filter(None, inputList))
    dict = {}
    for item in inputList:
        item = item.split(":")
        dict[item[0]] = float(item[1])
    return dict

class invoiceData:
    def __init__(self, invoice):
        self.invoice = invoice
        self.data = self.extractData()
    # retreive "Invoice Number","Total Quantity","Subtotal","Tax","Total"from DOC
    def extractData(self): #return a rowData list
        invoice = Document(self.invoice)
        #mine data from the word file paragraphs
        for paragraph in invoice.paragraphs:
            if paragraph.text.startswith("PRODUCTS"):
                productList = paragraph.text.split("\n")
                # remove Headers and space
                productList.remove('PRODUCTS')
                productDict = massageListToDict(productList)
            elif paragraph.text.startswith("SUBTOTAL"):
                monetaryList = paragraph.text.split("\n")
                monetaryDict = massageListToDict(monetaryList)
            else:
                invoiceNumber = []
                invoiceNumber.append(paragraph.text)
                invoiceNumber = ''.join(invoiceNumber)
        # put the data in a rowData List
        rowData = []
        rowData.append(invoiceNumber)
        rowData.append(int(sum(productDict.values())))
        rowData.append(monetaryDict['SUBTOTAL'])
        rowData.append(monetaryDict['TAX'])
        rowData.append(monetaryDict['TOTAL'])
        return rowData

class invoiceReport:
    def __init__(self,directory):
        self.directory=os.fsencode(directory)
        self.invoices = self.get_invoices()
        self.data = [invoice.data for invoice in self.invoices]

    def get_invoices(self):
        # directory = os.fsencode("Invoices")
        sheetData = []
        # loop through a directory and get all word document to document objects
        for file in os.listdir(self.directory):
            fileName = os.fsencode(file)
            if os.fsdecode(fileName).endswith(".docx") and not(os.fsdecode(fileName).startswith("~")):
                fullPath = os.fsdecode(os.path.join(self.directory,fileName))
                rowdata = invoiceData(fullPath)
                sheetData.append(rowdata)
        return sheetData
    
    def generateReport(self):
        workbook = Workbook()
        worksheet = workbook.active
        headerRow = ["Invoice Number","Total Quantity","Subtotal","Tax","Total"]
        worksheet.append(headerRow)
        for row in self.data:
            worksheet.append(row)
        workbook.save("Output_report-oop.xlsx")

def main():
    report = invoiceReport("Source")
    report.generateReport()

main()


